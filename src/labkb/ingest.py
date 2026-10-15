import io
import re
import zipfile
from pathlib import PurePosixPath

from docx import Document
from pypdf import PdfReader


class IngestError(ValueError):
    pass


def safe_name(filename: str) -> str:
    name = PurePosixPath(filename.replace("\\", "/")).name
    name = re.sub(r"[\x00-\x1f\x7f]", "", name).strip()
    if not name or len(name) > 180:
        raise IngestError("文件名为空或过长（最多 180 字符）")
    return name


def extract(name: str, data: bytes, max_chars: int) -> list[tuple[str, str]]:
    suffix = PurePosixPath(name).suffix.lower()
    sections = []
    try:
        if suffix in {".md", ".txt"}:
            text = data.decode("utf-8-sig")
            if "\x00" in text:
                raise IngestError("文本含二进制字符，请使用 UTF-8 文本")
            # Paragraph locations refer to blank-line separated blocks, not PDF pages.
            sections = [
                (f"段落 {i}", part.strip())
                for i, part in enumerate(re.split(r"\n\s*\n", text.replace("\r\n", "\n")), 1)
                if part.strip()
            ]
        elif suffix == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise IngestError("不支持加密 PDF，请先解密")
            if len(reader.pages) > 300:
                raise IngestError("PDF 最多 300 页，请拆分后上传")
            total = 0
            stream_total = 0
            for i, page in enumerate(reader.pages, 1):
                # Cap decompressed page streams before text extraction.
                contents = page.get_contents()
                stream_size = len(contents.get_data()) if contents else 0
                stream_total += stream_size
                if stream_size > 8 * 1024 * 1024 or stream_total > 30 * 1024 * 1024:
                    raise IngestError("PDF 页面过于复杂，请先拆分或转为文本")
                text = page.extract_text() or ""
                total += len(text)
                if total > max_chars:
                    raise IngestError("提取文本过大，请拆分文件")
                if text.strip():
                    sections.append((f"第 {i} 页", text.strip()))
        elif suffix == ".docx":
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                    raise IngestError("DOCX 解压后过大，请拆分文件")
            document = Document(io.BytesIO(data))
            sections = [
                (f"段落 {i}", p.text.strip())
                for i, p in enumerate(document.paragraphs, 1)
                if p.text.strip()
            ]
            for i, table in enumerate(document.tables, 1):
                for j, row in enumerate(table.rows, 1):
                    sections.append((f"表格 {i} 行 {j}", " | ".join(c.text for c in row.cells)))
        else:
            raise IngestError("支持 .md、.txt、.pdf、.docx 文件")
    except IngestError:
        raise
    except UnicodeDecodeError as exc:
        raise IngestError("文本编码必须为 UTF-8，请转换后重试") from exc
    except Exception as exc:
        raise IngestError("文件无法解析，请检查格式或转换为文本") from exc
    if sum(len(t) for _, t in sections) > max_chars:
        raise IngestError("提取文本过大，请拆分文件")
    if not sections or not any(t.strip() for _, t in sections):
        raise IngestError("未提取到文字；扫描 PDF 需先 OCR，本框架不内置 OCR")
    return sections


def chunk_sections(sections: list[tuple[str, str]], size: int, overlap: int) -> list[dict]:
    chunks = []
    for location, text in sections:
        start = 0
        while start < len(text):
            end = min(start + size, len(text))
            chunks.append(
                {"text": text[start:end], "location": f"{location} · 字符 {start + 1}–{end}"}
            )
            if end == len(text):
                break
            start = end - overlap
    return chunks
