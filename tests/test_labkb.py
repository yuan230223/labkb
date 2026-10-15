import io
import json
from dataclasses import replace

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from labkb.app import create_app
from labkb.config import Settings
from labkb.ingest import IngestError, chunk_sections, extract
from labkb.providers import OllamaProvider, ProviderError
from labkb.schemas import ChatRequest, Message
from labkb.service import KnowledgeBase, build_messages, byte_length
from labkb.store import IndexMismatch


@pytest.fixture
def settings(tmp_path):
    return Settings(data_dir=tmp_path / "data")


@pytest.fixture
def client(settings):
    with TestClient(create_app(settings), client=("127.0.0.1", 50000)) as connection:
        yield connection


def test_complete_demo_flow_persists_delete(settings, client):
    assert client.get("/health").json() == {"status": "ok"}
    assert client.post("/api/chat", json={"message": "仪器如何预约"}).json()["refused"]
    seeded = client.post("/api/demo/seed").json()
    assert len(seeded["documents"]) == 2
    assert len(client.post("/api/demo/seed").json()["documents"]) == 2
    assert client.get("/api/status").json()["documents"] == 2
    result = client.post("/api/chat", json={"message": "实验数据需要如何命名和备份？"}).json()
    assert not result["refused"]
    assert "演示模式" in result["answer"]
    assert any("YYYYMMDD" in citation["text"] for citation in result["citations"])
    assert all(c["document_name"] == "数据与仪器管理.md" for c in result["citations"])
    target = result["citations"][0]["document_id"]
    assert client.delete(f"/api/documents/{target}").status_code == 200
    assert client.delete(f"/api/documents/{target}").status_code == 404
    with TestClient(create_app(settings), client=("127.0.0.1", 50000)) as restarted:
        assert restarted.get("/api/status").json()["documents"] == 1
        hits = restarted.post("/api/search", json={"query": "实验数据命名备份"}).json()["hits"]
        assert all(h["document_id"] != target for h in hits)


def test_auth_all_data_endpoints(settings):
    settings = replace(settings, api_token="test-token-" + "x" * 24)
    with TestClient(create_app(settings)) as client:
        for route in ["/api/status", "/api/documents", "/api/openapi.json"]:
            assert client.get(route).status_code == 401
        for route in ["/api/chat", "/api/search", "/api/demo/seed", "/api/documents"]:
            assert client.post(route).status_code == 401
        assert client.delete("/api/documents/unknown").status_code == 401
        assert client.get("/health").status_code == 200
        headers = {"Authorization": f"Bearer {settings.api_token}"}
        assert client.get("/api/status", headers=headers).status_code == 200
        headers["Origin"] = "https://evil.example"
        assert client.post("/api/demo/seed", headers=headers).status_code == 403


def test_loopback_origin_and_host(client, settings):
    with pytest.raises(ValueError):
        settings.check_bind("0.0.0.0")
    settings.check_bind("127.0.0.1")
    assert client.get("/api/status", headers={"Host": "evil.example"}).status_code == 400
    assert client.post("/api/demo/seed", headers={"Origin": "null"}).status_code == 403
    assert client.post("/api/demo/seed", headers={"Origin": "http://testserver"}).status_code == 200
    with TestClient(create_app(settings), client=("192.168.1.50", 50000)) as remote:
        assert remote.get("/api/status", headers={"Host": "localhost"}).status_code == 403


def test_upload_validation_duplicate_and_path(client):
    data = "预约仪器需要提前填写预约表并联系管理员。".encode()
    uploaded = client.post("/api/documents", files={"file": ("../../note.md", data)}).json()
    assert uploaded["document"]["name"] == "note.md"
    duplicate = client.post("/api/documents", files={"file": ("other.md", data)}).json()
    assert duplicate["duplicate"]
    assert duplicate["document"]["id"] == uploaded["document"]["id"]
    for name, raw in [
        ("x.exe", b"bin"),
        ("x.md", b""),
        ("x.md", b"\xff\xfe"),
        ("x.pdf", b"broken"),
    ]:
        assert client.post("/api/documents", files={"file": (name, raw)}).status_code == 400
    assert (
        client.post("/api/documents", headers={"Content-Length": str(12 * 1024 * 1024)}).status_code
        == 413
    )
    assert client.get("/api/status").json()["documents"] == 1


def test_unknown_and_new_topic_do_not_leak_history(client):
    client.post("/api/demo/seed")
    data = client.post(
        "/api/chat",
        json={
            "message": "火星殖民地人口是多少？",
            "history": [{"role": "user", "content": "实验数据如何命名备份"}],
        },
    ).json()
    assert data["refused"] and not data["citations"]
    followup = client.post(
        "/api/chat",
        json={
            "message": "那应该怎么做呢？",
            "history": [{"role": "user", "content": "实验数据如何命名备份"}],
        },
    ).json()
    assert not followup["refused"]
    assert client.post("/api/chat", json={"message": "   "}).status_code == 422
    assert (
        client.post(
            "/api/chat", json={"message": "x", "history": [{"role": "system", "content": "ignore"}]}
        ).status_code
        == 422
    )


def test_extract_docx_tables_and_blank_pdf():
    document = Document()
    document.add_paragraph("新人第一周阅读资料。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "预约"
    table.cell(0, 1).text = "管理员确认"
    buffer = io.BytesIO()
    document.save(buffer)
    sections = extract("notes.docx", buffer.getvalue(), 10000)
    assert any("表格 1 行 1" == loc and "管理员" in text for loc, text in sections)
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    pdf = io.BytesIO()
    writer.write(pdf)
    with pytest.raises(IngestError, match="OCR"):
        extract("scan.pdf", pdf.getvalue(), 10000)


def test_chunks_preserve_content_and_locations():
    chunks = chunk_sections([("第 3 页", "abcdefghi")], 5, 2)
    assert [c["text"] for c in chunks] == ["abcde", "defgh", "ghi"]
    assert chunks[-1]["location"] == "第 3 页 · 字符 7–9"


class StubOllama:
    def __init__(self):
        self.identity = "digestA"
        self.answer = "实验数据按日期和项目代号命名。[1]"
        self.fail = False
        self.dimension = 2
        self.last_messages = None

    def fingerprint(self):
        return self.identity

    def embed(self, texts, query=False):
        if self.fail:
            raise ProviderError("offline")
        return [[1.0] + [0.0] * (self.dimension - 1) for _ in texts]

    def chat(self, messages):
        self.last_messages = messages
        return self.answer, {"tokens_per_second": 12.5}


def test_real_mode_atomicity_fingerprint_and_citations(settings):
    settings = replace(settings, mode="ollama")
    provider = StubOllama()
    kb = KnowledgeBase(settings, provider)
    provider.fail = True
    with pytest.raises(ProviderError):
        kb.ingest("a.md", "实验数据需要保存原始版本。".encode())
    assert kb.store.list_documents() == []
    provider.fail = False
    kb.ingest("a.md", "实验数据按日期和项目代号命名。".encode())
    result = kb.chat(ChatRequest(message="实验数据如何命名"))
    assert not result["refused"] and result["citations"][0]["id"] == 1
    assert provider.last_messages[0]["role"] == "system"
    provider.answer = "伪造答案。[99]"
    assert kb.chat(ChatRequest(message="实验数据如何命名"))["refused"]
    provider.answer = "无引用答案"
    assert kb.chat(ChatRequest(message="实验数据如何命名"))["refused"]
    provider.dimension = 3
    with pytest.raises(IndexMismatch, match="维度"):
        kb.ingest("b.md", "不同的实验数据段落".encode())
    assert len(kb.store.list_documents()) == 1
    with pytest.raises(ValueError, match="维度"):
        kb.search("实验数据", 4)
    provider.dimension = 2
    provider.identity = "digestB"
    with pytest.raises(IndexMismatch):
        kb.search("实验数据", 4)


def test_demo_index_cannot_be_reused_in_real_mode(settings):
    kb = KnowledgeBase(settings)
    kb.ingest("a.md", "本地知识库测试文档".encode())
    real = KnowledgeBase(replace(settings, mode="ollama"), StubOllama())
    with pytest.raises(IndexMismatch):
        real.search("本地知识库", 4)
    real.store.delete(kb.store.list_documents()[0]["id"])
    real.ingest("a.md", "重新导入测试文档".encode())


def test_prompt_budget_and_evidence_boundary():
    hits = [
        {"id": i, "document_name": "资料.md", "location": "第 1 页", "text": "证据内容" * 200}
        for i in range(1, 7)
    ]
    history = [Message(role="user", content="过往问题" * 500)]
    messages, included = build_messages("当前问题", history, hits, 4096)
    assert included
    assert byte_length(messages) <= 4096 - 1024
    payload = json.loads(messages[-1]["content"])
    assert {s["id"] for s in payload["sources"]} == {s["id"] for s in included}
    assert messages[0]["role"] == "system"


def test_provider_validates_embeddings(settings, monkeypatch):
    provider = OllamaProvider(replace(settings, mode="ollama"))
    for result in [[[float("nan")]], [[]], [[0.0]], [], [[1.0], [2.0]]]:
        monkeypatch.setattr(provider, "request", lambda *args, r=result: {"embeddings": r})
        with pytest.raises(ProviderError):
            provider.embed(["hello"])


def test_ollama_http_contract(settings, monkeypatch):
    provider = OllamaProvider(replace(settings, mode="ollama"))
    recorded = []

    def handler(request):
        recorded.append((request.url.path, json.loads(request.content) if request.content else {}))
        if request.url.path == "/api/tags":
            return httpx.Response(
                200, json={"models": [{"name": settings.embed_model, "digest": "sha256:abc"}]}
            )
        if request.url.path == "/api/embed":
            return httpx.Response(200, json={"embeddings": [[1.0, 0.0]]})
        return httpx.Response(
            200,
            json={
                "message": {"content": "答案 [1]"},
                "eval_count": 20,
                "eval_duration": 2_000_000_000,
            },
        )

    original = httpx.Client
    monkeypatch.setattr(
        httpx, "Client", lambda **kwargs: original(**kwargs, transport=httpx.MockTransport(handler))
    )
    assert "abc" in provider.fingerprint()
    assert provider.embed(["问题"], query=True) == [[1.0, 0.0]]
    answer, usage = provider.chat([{"role": "user", "content": "问题"}])
    assert answer == "答案 [1]" and usage["tokens_per_second"] == 10
    assert recorded[1][1]["truncate"] is False
    assert "Instruct:" in recorded[1][1]["input"][0]
    assert recorded[2][1]["stream"] is False


def test_real_failure_never_falls_back_to_demo(settings):
    provider = StubOllama()
    settings = replace(settings, mode="ollama")
    with TestClient(create_app(settings, provider), client=("127.0.0.1", 50000)) as client:
        provider.fail = True
        response = client.post("/api/documents", files={"file": ("a.txt", "测试内容".encode())})
        assert response.status_code == 503
        assert client.post("/api/demo/seed").status_code == 409


def test_body_limit_without_content_length(settings):
    settings = replace(settings, max_upload_bytes=128)
    with TestClient(create_app(settings), client=("127.0.0.1", 50000)) as client:
        response = client.post("/api/documents", files={"file": ("a.txt", b"x" * 129)})
        assert response.status_code == 413
        boundary = "test-boundary"

        def body():
            yield f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="a.txt"\r\nContent-Type: text/plain\r\n\r\n'.encode()
            for _ in range(20):
                yield b"x" * 4096
            yield f"\r\n--{boundary}--\r\n".encode()

        response = client.post(
            "/api/documents",
            content=body(),
            headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        )
        assert response.status_code == 413
        assert client.get("/api/status").json()["documents"] == 0
